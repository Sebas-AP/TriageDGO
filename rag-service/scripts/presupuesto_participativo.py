"""Reglas auditables para presupuesto participativo mensual de Durango.

El puntaje base reproduce la hoja Calculadora_Presupuesto del archivo fuente:
  S_base = 0.40 * (poblacion / total_poblacion)
         + 0.60 * (rezago / total_rezago)

La prioridad de recaudación no sustituye el rezago: ajusta ese puntaje con la
participación mensual de predial y se normaliza para conservar el techo exacto.
"""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

import openpyxl

MONEY = Decimal("0.01")


def normalize_key(value: object) -> str:
    return "".join(str(value or "").upper().split())


def read_zones(workbook_path: Path) -> tuple[Decimal, Decimal, Decimal, list[dict[str, Any]]]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=False, read_only=True)
    sheet = workbook["Calculadora_Presupuesto"]
    monthly_cap = Decimal(str(sheet["C6"].value))
    population_weight = Decimal(str(sheet["C7"].value))
    deprivation_weight = Decimal("1") - population_weight
    zones: list[dict[str, Any]] = []
    for row in range(12, sheet.max_row + 1):
        key = sheet.cell(row, 1).value
        population, deprivation = sheet.cell(row, 4).value, sheet.cell(row, 5).value
        if key is None or population is None or deprivation is None: continue
        try:
            population_value, deprivation_value = Decimal(str(population)), Decimal(str(deprivation))
        except Exception:
            # Omite filas de totales y leyendas como "TOTAL MUNICIPAL".
            continue
        zones.append({"clave_catastral": str(key), "clave_normalizada": normalize_key(key), "colonia": str(sheet.cell(row, 2).value), "sector": str(sheet.cell(row, 3).value), "poblacion": population_value, "rezago_social": deprivation_value})
    if not zones or monthly_cap <= 0: raise ValueError("El libro no contiene zonas válidas o un techo presupuestal positivo")
    return monthly_cap, population_weight, deprivation_weight, zones


def parse_month(value: str) -> tuple[date, date]:
    try: start = datetime.strptime(value, "%Y-%m").date().replace(day=1)
    except ValueError as error: raise ValueError("El mes debe tener formato AAAA-MM") from error
    end = date(start.year + (start.month == 12), 1 if start.month == 12 else start.month + 1, 1)
    return start, end


def payment_month(payment: dict[str, Any]) -> date | None:
    raw = payment.get("fecha_pago") or payment.get("fechaPago") or payment.get("actualizadoEn")
    if not raw: return None
    try: return datetime.fromisoformat(str(raw).replace("Z", "+00:00")).date()
    except ValueError: return None


def paid_amount(payment: dict[str, Any]) -> Decimal:
    raw = payment.get("monto_pagado", payment.get("montoPagado", 0))
    try: return Decimal(str(raw))
    except Exception: return Decimal("0")


def payments_by_zone(payments: Iterable[dict[str, Any]], zones: list[dict[str, Any]], month: str) -> dict[str, Decimal]:
    start, end = parse_month(month)
    known = {zone["clave_normalizada"]: zone["clave_catastral"] for zone in zones}
    totals: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for payment in payments:
        paid_at = payment_month(payment)
        status = str(payment.get("estatus", "pagado")).lower()
        if not paid_at or not start <= paid_at < end or status not in {"pagado", "aplicado", "completado"}: continue
        key = normalize_key(payment.get("clave_zona") or payment.get("claveZona") or payment.get("clave_catastral") or payment.get("claveCatastral"))
        if key in known: totals[known[key]] += max(Decimal("0"), paid_amount(payment))
    return totals


def calculate(workbook_path: Path, payments: Iterable[dict[str, Any]], month: str, monthly_cap: Decimal | None = None, tax_weight: Decimal = Decimal("0.15")) -> dict[str, Any]:
    if not Decimal("0") <= tax_weight <= Decimal("1"): raise ValueError("tax_weight debe estar entre 0 y 1")
    source_cap, population_weight, deprivation_weight, zones = read_zones(workbook_path)
    cap = monthly_cap if monthly_cap is not None else source_cap
    if cap <= 0: raise ValueError("El techo mensual debe ser positivo")
    paid = payments_by_zone(payments, zones, month)
    total_population = sum((zone["poblacion"] for zone in zones), Decimal("0"))
    total_deprivation = sum((zone["rezago_social"] for zone in zones), Decimal("0"))
    total_paid = sum(paid.values(), Decimal("0"))
    provisional: list[dict[str, Any]] = []
    for zone in zones:
        population_factor = zone["poblacion"] / total_population
        deprivation_factor = zone["rezago_social"] / total_deprivation
        base_score = population_weight * population_factor + deprivation_weight * deprivation_factor
        paid_amount = paid[zone["clave_catastral"]]
        paid_factor = paid_amount / total_paid if total_paid else Decimal("0")
        # Con cero pagos se conserva la fórmula exacta del Excel.
        adjusted_score = base_score if not total_paid else base_score * ((Decimal("1") - tax_weight) + tax_weight * paid_factor * len(zones))
        provisional.append({**zone, "factor_poblacion": population_factor, "factor_rezago": deprivation_factor, "puntaje_base": base_score, "predial_pagado": paid_amount, "participacion_predial": paid_factor, "puntaje_ajustado": adjusted_score})
    score_total = sum((row["puntaje_ajustado"] for row in provisional), Decimal("0"))
    if not score_total: raise ValueError("No fue posible calcular puntajes")
    remaining = cap
    for row in provisional:
        allocation = (cap * row["puntaje_ajustado"] / score_total).quantize(MONEY, rounding=ROUND_HALF_UP)
        row["presupuesto_asignado"] = allocation
        remaining -= allocation
    # El ajuste de centavos mantiene la igualdad exacta entre techo y asignaciones.
    provisional[0]["presupuesto_asignado"] += remaining
    provisional.sort(key=lambda row: (-row["predial_pagado"], row["colonia"]))
    for position, row in enumerate(provisional, start=1): row["prioridad_recaudacion"] = position
    provisional.sort(key=lambda row: row["clave_catastral"])
    for row in provisional:
        for field in ("poblacion", "rezago_social", "factor_poblacion", "factor_rezago", "puntaje_base", "predial_pagado", "participacion_predial", "puntaje_ajustado", "presupuesto_asignado"): row[field] = str(row[field])
        row.pop("clave_normalizada", None)
    return {"mes": month, "techo_mensual": str(cap), "formula_base": {"peso_poblacion": str(population_weight), "peso_rezago": str(deprivation_weight)}, "peso_prioridad_predial": str(tax_weight), "total_predial_pagado": str(total_paid), "zonas": provisional}


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def firestore_client() -> Any:
    import firebase_admin
    from firebase_admin import credentials, firestore
    if not firebase_admin._apps:
        key = __import__("os").getenv("GOOGLE_APPLICATION_CREDENTIALS")
        firebase_admin.initialize_app(credentials.Certificate(key) if key else credentials.ApplicationDefault())
    return firestore.client()


def read_firestore_payments() -> list[dict[str, Any]]:
    """Lee la colección transaccional; `predial` sólo guarda el estado actual."""
    return [doc.to_dict() for doc in firestore_client().collection("pagos_predial").stream()]


def previous_month(reference: date | None = None) -> str:
    ref = reference or datetime.now(timezone.utc).date()
    return f"{ref.year - (ref.month == 1):04d}-{12 if ref.month == 1 else ref.month - 1:02d}"
